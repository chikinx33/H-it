#!/usr/bin/env python3
"""포트폴리오 덱 검사: 빈 placeholder, 이미지 비율 왜곡.
사용: python3 tools/check_deck.py <deck.pptx> [...]
"""
import io, sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

TOL = 0.02  # 비율 허용 오차 2%

def check(path):
    prs = Presentation(path)
    bad = 0
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
                print(f"  [빈 텍스트 창] {i}번 슬라이드: placeholder '{sh.name}'")
                bad += 1
            if sh.shape_type == 13:  # PICTURE
                im = Image.open(io.BytesIO(sh.image.blob))
                iw, ih = im.size
                l, r, t, b = sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom
                vis_r = (iw * (1 - l - r)) / (ih * (1 - t - b))
                box_r = sh.width / sh.height
                if abs(box_r / vis_r - 1) > TOL:
                    print(f"  [이미지 왜곡] {i}번 슬라이드: '{sh.name}' 원본 {vis_r:.3f} / 배치 {box_r:.3f}")
                    bad += 1
    print(f"{path}: {'통과' if bad == 0 else f'문제 {bad}건'}")
    return bad

if __name__ == "__main__":
    sys.exit(1 if sum(check(p) for p in sys.argv[1:]) else 0)
